from __future__ import annotations

import csv
import io
import re
import unicodedata
from collections import Counter
from dataclasses import dataclass
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from typing import Literal
from zipfile import BadZipFile
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field, model_validator

UNIDADES_INTERVALO = (5, 15, 60)
LIMITE_REGISTROS = 500_000
CAMPOS = frozenset({"ts", "kwh_ativo", "kvarh_indutivo", "kvarh_capacitivo", "kw"})
ALIASES: dict[str, frozenset[str]] = {
    "ts": frozenset(
        {
            "data hora",
            "datahora",
            "timestamp",
            "data e hora",
            "data leitura",
            "horario",
            "date time",
            "datetime",
        }
    ),
    "kwh_ativo": frozenset(
        {
            "kwh ativo",
            "energia ativa kwh",
            "energia ativa",
            "consumo ativo kwh",
            "consumo ativo",
            "kwh",
            "active energy kwh",
        }
    ),
    "kvarh_indutivo": frozenset(
        {
            "kvarh indutivo",
            "energia reativa indutiva kvarh",
            "reativo indutivo",
            "reativo indutivo kvarh",
            "kvarh ind",
        }
    ),
    "kvarh_capacitivo": frozenset(
        {
            "kvarh capacitivo",
            "energia reativa capacitiva kvarh",
            "reativo capacitivo",
            "reativo capacitivo kvarh",
            "kvarh cap",
        }
    ),
    "kw": frozenset({"kw", "potencia kw", "demanda kw", "potencia media kw", "active power kw"}),
}


class PerfilImportacao(BaseModel):
    model_config = ConfigDict(extra="forbid")

    separador: str | None = Field(None, min_length=1, max_length=1)
    decimal: Literal["auto", ".", ","] = "auto"
    codificacao: str | None = None
    fuso: str = "America/Sao_Paulo"
    formato_data_hora: str | None = None
    colunas: dict[str, str] | None = None
    unidade: Literal["KWH_INTERVALO", "KW_MEDIO"] = "KWH_INTERVALO"
    intervalo_min: Literal[5, 15, 60] | None = None
    fonte: Literal["DISTRIBUIDORA", "MEDIDOR_PARALELO"] = "DISTRIBUIDORA"

    @model_validator(mode="after")
    def validar_colunas(self) -> PerfilImportacao:
        if self.colunas is not None and (
            not set(self.colunas).issubset(CAMPOS)
            or "ts" not in self.colunas
            or not ({"kwh_ativo", "kw"} & set(self.colunas))
        ):
            raise ValueError("Mapeie ts e pelo menos um campo entre kwh_ativo e kw.")
        if self.unidade == "KW_MEDIO" and (self.colunas is None or "kw" not in self.colunas):
            raise ValueError("A unidade KW_MEDIO exige o mapeamento da coluna kw.")
        if self.unidade == "KW_MEDIO" and self.intervalo_min is None:
            raise ValueError("A unidade KW_MEDIO exige intervalo_min no perfil.")
        if (
            self.unidade == "KWH_INTERVALO"
            and self.colunas is not None
            and "kwh_ativo" not in self.colunas
        ):
            raise ValueError("A unidade KWH_INTERVALO exige o mapeamento da coluna kwh_ativo.")
        return self


@dataclass(frozen=True, slots=True)
class RegistroMedicao:
    ts_inicio: datetime
    intervalo_min: int
    kwh_ativo: Decimal
    kvarh_indutivo: Decimal | None
    kvarh_capacitivo: Decimal | None


@dataclass(frozen=True, slots=True)
class ResultadoImportacao:
    perfil: PerfilImportacao
    registros: tuple[RegistroMedicao, ...]
    colunas_encontradas: tuple[str, ...]
    lacunas: int
    duplicatas: int
    valores_negativos: int
    intervalos_irregulares: int
    completude_pct: Decimal
    alertas: tuple[str, ...]
    erros_bloqueantes: tuple[str, ...]


class ErroImportacaoMedicao(ValueError):
    def __init__(
        self,
        detail: str,
        *,
        colunas_encontradas: tuple[str, ...] = (),
    ) -> None:
        super().__init__(detail)
        self.colunas_encontradas = colunas_encontradas


def parsear_memoria(
    conteudo: bytes,
    *,
    perfil: PerfilImportacao | None = None,
) -> ResultadoImportacao:
    """Interpreta CSV/XLSX, normaliza horários a UTC e valida a série (§11.1)."""
    if not conteudo:
        raise ErroImportacaoMedicao("Arquivo vazio.")
    perfil = perfil or PerfilImportacao(separador=None)
    try:
        ZoneInfo(perfil.fuso)
    except (ValueError, ZoneInfoNotFoundError) as error:
        raise ErroImportacaoMedicao("Fuso horário inválido no perfil.") from error

    if conteudo.startswith(b"PK\x03\x04"):
        headers, rows = _ler_xlsx(conteudo)
    else:
        headers, rows, detected_separator = _ler_csv(conteudo, perfil)
        if perfil.separador is None:
            perfil = perfil.model_copy(update={"separador": detected_separator})
    if not headers:
        raise ErroImportacaoMedicao("O arquivo não contém cabeçalho.")

    columns = _mapear_colunas(headers, perfil)
    if columns is None:
        raise ErroImportacaoMedicao(
            "Não foi possível detectar o cabeçalho; informe um mapeamento de colunas.",
            colunas_encontradas=tuple(headers),
        )
    if "kwh_ativo" not in columns and "kw" not in columns:
        raise ErroImportacaoMedicao(
            "Mapeie ao menos a energia ativa (kWh) ou potência média (kW).",
            colunas_encontradas=tuple(headers),
        )
    if perfil.unidade == "KW_MEDIO" and "kw" not in columns:
        raise ErroImportacaoMedicao(
            "A unidade KW_MEDIO exige uma coluna de potência média (kW).",
            colunas_encontradas=tuple(headers),
        )
    if perfil.unidade == "KWH_INTERVALO" and "kwh_ativo" not in columns:
        raise ErroImportacaoMedicao(
            "A unidade KWH_INTERVALO exige uma coluna de energia ativa (kWh).",
            colunas_encontradas=tuple(headers),
        )

    decimal_sep = _decimal_separator(perfil, rows, columns)
    records: list[RegistroMedicao] = []
    seen: set[datetime] = set()
    observed_timestamps: set[datetime] = set()
    duplicates = 0
    negatives = 0
    row_errors: list[int] = []
    for row_number, row in enumerate(rows, start=2):
        if not any(value not in (None, "") for value in row.values()):
            continue
        try:
            timestamp = _parse_timestamp(
                row.get(columns["ts"]),
                formato=perfil.formato_data_hora,
                fuso=perfil.fuso,
            )
            if timestamp in seen:
                duplicates += 1
                continue
            seen.add(timestamp)
            observed_timestamps.add(timestamp)
            active_kwh = (
                _parse_decimal(row.get(columns["kw"]), decimal_sep)
                * Decimal(perfil.intervalo_min or 0)
                / Decimal(60)
                if perfil.unidade == "KW_MEDIO"
                else _parse_decimal(row.get(columns["kwh_ativo"]), decimal_sep)
            )
            inductive = (
                _parse_optional_decimal(row.get(columns["kvarh_indutivo"]), decimal_sep)
                if "kvarh_indutivo" in columns
                else None
            )
            capacitive = (
                _parse_optional_decimal(row.get(columns["kvarh_capacitivo"]), decimal_sep)
                if "kvarh_capacitivo" in columns
                else None
            )
            values = (active_kwh, inductive, capacitive)
            if any(value is not None and value < 0 for value in values):
                negatives += 1
                continue
            records.append(
                RegistroMedicao(
                    ts_inicio=timestamp,
                    intervalo_min=0,
                    kwh_ativo=active_kwh,
                    kvarh_indutivo=inductive,
                    kvarh_capacitivo=capacitive,
                )
            )
        except (InvalidOperation, ValueError, TypeError, OverflowError):
            row_errors.append(row_number)

    if row_errors:
        raise ErroImportacaoMedicao(
            f"Valores ou data/hora inválidos nas linhas: {', '.join(map(str, row_errors[:20]))}."
        )
    if not records:
        raise ErroImportacaoMedicao("Nenhuma medição válida foi encontrada.")

    records.sort(key=lambda item: item.ts_inicio)
    interval = perfil.intervalo_min or _inferir_intervalo(sorted(observed_timestamps))
    records = [
        RegistroMedicao(
            ts_inicio=item.ts_inicio,
            intervalo_min=interval,
            kwh_ativo=item.kwh_ativo,
            kvarh_indutivo=item.kvarh_indutivo,
            kvarh_capacitivo=item.kvarh_capacitivo,
        )
        for item in records
    ]
    irregular = 0
    gaps = 0
    step = timedelta(minutes=interval)
    for previous, current in zip(records, records[1:], strict=False):
        delta = current.ts_inicio - previous.ts_inicio
        if delta <= timedelta(0) or delta.total_seconds() % step.total_seconds():
            irregular += 1
        elif delta > step:
            gaps += int(delta / step) - 1
    expected = int((records[-1].ts_inicio - records[0].ts_inicio) / step) + 1
    completeness = (Decimal(len(records)) * Decimal(100) / Decimal(expected)).quantize(
        Decimal("0.01")
    )
    blockers: list[str] = []
    alerts: list[str] = []
    if duplicates:
        blockers.append("DUPLICATAS")
    if negatives:
        blockers.append("VALORES_NEGATIVOS")
    if irregular:
        blockers.append("INTERVALO_IRREGULAR")
    if completeness < Decimal(95):
        alerts.append("DADOS_INSUFICIENTES")
    effective_profile = perfil.model_copy(update={"intervalo_min": interval})
    return ResultadoImportacao(
        perfil=effective_profile,
        registros=tuple(records),
        colunas_encontradas=tuple(headers),
        lacunas=gaps,
        duplicatas=duplicates,
        valores_negativos=negatives,
        intervalos_irregulares=irregular,
        completude_pct=completeness,
        alertas=tuple(alerts),
        erros_bloqueantes=tuple(blockers),
    )


def _ler_csv(
    conteudo: bytes, perfil: PerfilImportacao
) -> tuple[list[str], list[dict[str, object]], str]:
    decoded: str | None = None
    encodings = (perfil.codificacao,) if perfil.codificacao else ("utf-8-sig", "cp1252")
    for encoding in encodings:
        if encoding is None:
            continue
        try:
            decoded = conteudo.decode(encoding)
            break
        except (LookupError, UnicodeDecodeError):
            continue
    if decoded is None:
        raise ErroImportacaoMedicao("A codificação do arquivo não é compatível com o perfil.")
    sample = decoded[:8192]
    separator = perfil.separador
    if separator is None:
        try:
            separator = csv.Sniffer().sniff(sample, delimiters=";,\t|").delimiter
        except csv.Error as error:
            raise ErroImportacaoMedicao("Não foi possível detectar o separador CSV.") from error
    reader = csv.reader(io.StringIO(decoded), delimiter=separator)
    try:
        headers = [header.strip() for header in next(reader)]
    except StopIteration as error:
        raise ErroImportacaoMedicao("O arquivo não contém cabeçalho.") from error
    rows: list[dict[str, object]] = []
    for values in reader:
        if len(values) != len(headers):
            raise ErroImportacaoMedicao("Há linhas CSV com quantidade de colunas inconsistente.")
        if len(rows) >= LIMITE_REGISTROS:
            raise ErroImportacaoMedicao("O arquivo excede o limite de medições por importação.")
        rows.append(dict(zip(headers, values, strict=True)))
    return headers, rows, separator


def _ler_xlsx(conteudo: bytes) -> tuple[list[str], list[dict[str, object]]]:
    try:
        workbook = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
        try:
            worksheet = workbook.active
            if worksheet is None:
                return [], []
            if worksheet.max_row is not None and worksheet.max_row > LIMITE_REGISTROS + 1:
                raise ErroImportacaoMedicao("O arquivo excede o limite de medições por importação.")
            rows_iterator = worksheet.iter_rows(values_only=True)
            raw_headers = next(rows_iterator, None)
            if raw_headers is None:
                return [], []
            headers = [str(value).strip() if value is not None else "" for value in raw_headers]
            if not any(headers):
                return [], []
            rows: list[dict[str, object]] = []
            for values in rows_iterator:
                if not any(value not in (None, "") for value in values):
                    continue
                if len(rows) >= LIMITE_REGISTROS:
                    raise ErroImportacaoMedicao(
                        "O arquivo excede o limite de medições por importação."
                    )
                rows.append(dict(zip(headers, values, strict=True)))
            return headers, rows
        finally:
            workbook.close()
    except ErroImportacaoMedicao:
        raise
    except (BadZipFile, OSError, ValueError, KeyError, StopIteration) as error:
        raise ErroImportacaoMedicao("Não foi possível ler a planilha XLSX.") from error


def _mapear_colunas(headers: list[str], perfil: PerfilImportacao) -> dict[str, str] | None:
    if perfil.colunas is not None:
        if any(column not in headers for column in perfil.colunas.values()):
            raise ErroImportacaoMedicao(
                "O mapeamento referencia colunas ausentes no arquivo.",
                colunas_encontradas=tuple(headers),
            )
        return perfil.colunas
    normalized = {_normalizar(header): header for header in headers}
    mapped: dict[str, str] = {}
    for field, aliases in ALIASES.items():
        match = next((normalized[alias] for alias in aliases if alias in normalized), None)
        if match is not None:
            mapped[field] = match
    return mapped if "ts" in mapped else None


def _decimal_separator(
    perfil: PerfilImportacao,
    rows: list[dict[str, object]],
    columns: dict[str, str],
) -> str:
    if perfil.decimal != "auto":
        return perfil.decimal
    numeric_fields = [field for field in ("kwh_ativo", "kw") if field in columns]
    for row in rows:
        value = str(row.get(columns[numeric_fields[0]], "")).strip()
        if "," in value and "." in value:
            return "," if value.rfind(",") > value.rfind(".") else "."
        if "," in value:
            return ","
        if "." in value:
            return "."
    return "."


def _parse_timestamp(value: object, *, formato: str | None, fuso: str) -> datetime:
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime.combine(value, datetime.min.time())
    elif isinstance(value, str):
        normalized = value.strip()
        if not normalized:
            raise ValueError("Data/hora vazia.")
        try:
            parsed = datetime.fromisoformat(normalized.replace("Z", "+00:00"))
        except ValueError:
            formats = (
                (formato,)
                if formato
                else (
                    "%d/%m/%Y %H:%M:%S",
                    "%d/%m/%Y %H:%M",
                    "%Y-%m-%d %H:%M:%S",
                    "%Y-%m-%d %H:%M",
                    "%d/%m/%Y",
                )
            )
            parsed = None
            for candidate_format in formats:
                if candidate_format is None:
                    continue
                try:
                    parsed = datetime.strptime(normalized, candidate_format)
                    break
                except ValueError:
                    continue
            if parsed is None:
                raise ValueError("Formato de data/hora inválido.") from None
    else:
        raise ValueError("Data/hora inválida.")
    if parsed is None:
        raise ValueError("Formato de data/hora inválido.")
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=ZoneInfo(fuso))
    return parsed.astimezone(UTC)


def _parse_optional_decimal(value: object, separator: str) -> Decimal | None:
    if value is None or str(value).strip() == "":
        return None
    return _parse_decimal(value, separator)


def _parse_decimal(value: object, separator: str) -> Decimal:
    if value is None:
        raise ValueError("Valor numérico ausente.")
    text = str(value).strip().replace("\u00a0", "").replace(" ", "")
    if separator == ",":
        text = text.replace(".", "").replace(",", ".")
    else:
        text = text.replace(",", "")
    result = Decimal(text)
    if not result.is_finite():
        raise ValueError("Valor numérico não finito.")
    return result


def _inferir_intervalo(timestamps: list[datetime]) -> int:
    differences = [
        int((right - left).total_seconds() // 60)
        for left, right in zip(timestamps, timestamps[1:], strict=False)
        if right > left
    ]
    candidates = [difference for difference in differences if difference in UNIDADES_INTERVALO]
    if candidates:
        return Counter(candidates).most_common(1)[0][0]
    raise ErroImportacaoMedicao(
        "Não foi possível inferir o intervalo; informe 5, 15 ou 60 minutos no perfil."
    )


def _normalizar(value: str) -> str:
    ascii_text = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", ascii_text.casefold())).strip()
